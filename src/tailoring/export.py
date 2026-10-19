"""Export a tailored CV as DOCX or PDF with a plain, ATS-friendly single-column layout.

Headings come from recognised section names, "- " lines become bullets, and the first line
(usually the name) becomes the title. Content is written as text, never as images.
"""

from __future__ import annotations

import io
from html import escape

import pymupdf as fitz
from docx import Document
from docx.shared import Pt

from src.cv.parsing import match_heading

_PDF_CSS = """
body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.35; color: #172033; }
h1 { font-size: 18pt; margin: 0 0 4pt 0; }
h2 { font-size: 12pt; margin: 12pt 0 4pt 0; padding-bottom: 2pt; border-bottom: 0.6pt solid #8A94A6;
     text-transform: uppercase; letter-spacing: 0.5pt; }
p { margin: 0 0 3pt 0; }
ul { margin: 0 0 4pt 0; padding-left: 14pt; }
li { margin: 0 0 2pt 0; }
"""


def _blocks(text: str) -> list[tuple[str, str]]:
    """(kind, text) with kind in title / heading / bullet / line / blank."""
    out: list[tuple[str, str]] = []
    first = True
    for raw in text.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            out.append(("blank", ""))
            continue
        if first:
            out.append(("title", line.strip()))
            first = False
        elif match_heading(line) is not None:
            out.append(("heading", line.strip().rstrip(":")))
        elif line.startswith("- "):
            out.append(("bullet", line[2:].strip()))
        else:
            out.append(("line", line.strip()))
    return out


def to_docx(text: str) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for kind, value in _blocks(text):
        if kind == "title":
            doc.add_heading(value, level=0)
        elif kind == "heading":
            doc.add_heading(value.title() if value.isupper() else value, level=1)
        elif kind == "bullet":
            doc.add_paragraph(value, style="List Bullet")
        elif kind == "line":
            doc.add_paragraph(value)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _html(text: str) -> str:
    parts: list[str] = []
    in_list = False
    for kind, value in _blocks(text):
        if kind != "bullet" and in_list:
            parts.append("</ul>")
            in_list = False
        if kind == "title":
            parts.append(f"<h1>{escape(value)}</h1>")
        elif kind == "heading":
            parts.append(f"<h2>{escape(value)}</h2>")
        elif kind == "bullet":
            if not in_list:
                parts.append("<ul>")
                in_list = True
            parts.append(f"<li>{escape(value)}</li>")
        elif kind == "line":
            parts.append(f"<p>{escape(value)}</p>")
    if in_list:
        parts.append("</ul>")
    return f"<body>{''.join(parts)}</body>"


def to_pdf(text: str) -> bytes:
    """A4 PDF with real, selectable text (laid out with PyMuPDF's Story)."""
    story = fitz.Story(html=_html(text), user_css=_PDF_CSS)
    buf = io.BytesIO()
    writer = fitz.DocumentWriter(buf)
    page = fitz.paper_rect("a4")
    area = page + (54, 54, -54, -54)  # 0.75-inch margins
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(area)
        story.draw(device)
        writer.end_page()
    writer.close()
    return buf.getvalue()
