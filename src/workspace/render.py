"""Generate the final CV as a PDF and an editable DOCX, then check both files.

Layout rules: a single column, real (selectable) text, standard headings, plain bullets and no
text boxes, images or tables — what ATS software reads most reliably. Multi-page CVs keep each
heading with the text that follows it and get page numbers. Every generated file is opened
again and checked: page count, every visible line present, reading order, nothing clipped or
overlapping. A file that fails a check is reported, never presented as a successful download.
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from html import escape

import pymupdf as fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from src.workspace.document import INLINE_KINDS, CVDocument, Section
from src.workspace.style import MIN_BODY_SIZE, MIN_MARGIN, PAGE_SIZES, FormatSettings

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_FONT_STACK = {"sans-serif": "sans-serif", "serif": "serif"}
_DOCX_FONT = {"sans-serif": "Calibri", "serif": "Cambria"}
_MAX_LAYOUT_PASSES = 16


class RenderError(Exception):
    """The document could not be generated; the message is safe to show."""


# --- shared helpers ------------------------------------------------------------------------------

def visible_sections(content: CVDocument) -> list[Section]:
    out = []
    for s in content.sections:
        if s.hidden:
            continue
        entries = [e for e in s.entries if not e.hidden and (e.header or any(not i.hidden for i in e.items))]
        if entries:
            out.append(s.model_copy(update={"entries": [e.model_copy(update={"items": [i for i in e.items
                                                                                    if not i.hidden]})
                                                        for e in entries]}))
    return out


def contact_line(content: CVDocument) -> str:
    c = content.contact
    return "  ·  ".join(x for x in (c.email, c.phone, c.location, *c.links) if x)


_TRANSLIT = str.maketrans({"Ł": "L", "ł": "l", "Ø": "O", "ø": "o", "Đ": "D", "đ": "d", "ß": "ss", "Æ": "AE",
                           "æ": "ae", "Œ": "OE", "œ": "oe", "Þ": "Th", "þ": "th", "ı": "i"})
_FOOTER_LINE = re.compile(r"(?im)^.*\bpage \d+ of \d+[ \t]*\n?")


def professional_filename(content: CVDocument, target_title: str | None, ext: str) -> str:
    """First_Last_Title_CV.pdf — ASCII only, safe on every operating system."""
    def clean(text: str) -> str:
        text = (text or "").translate(_TRANSLIT)
        text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
        return "_".join(w for w in re.sub(r"[^A-Za-z0-9]+", " ", text).split())

    parts = [p for p in (clean(content.contact.name)[:60], clean(target_title or "")[:60]) if p]
    return "_".join([*parts, "CV"]) + f".{ext}"


def _norm(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).replace("­", "")
    text = re.sub(r"[‐‑‒–—]", "-", text).replace("’", "'").replace("“", '"').replace("”", '"')
    return " ".join(text.split()).lower()


def expected_lines(content: CVDocument) -> list[str]:
    """Every piece of visible text the generated document must contain, in reading order."""
    lines = [x for x in (content.contact.name, content.contact.headline) if x]
    for s in visible_sections(content):
        lines.append(s.title)
        for e in s.entries:
            lines += [x for x in (e.title, e.organization, e.location, e.dates) if x]
            lines += [i.text for i in e.items]
    return lines


# --- PDF ------------------------------------------------------------------------------------

def _css(fmt: FormatSettings) -> str:
    family = _FONT_STACK[fmt.font_family]
    body, head, name, accent = fmt.body_size, fmt.heading_size, fmt.name_size, fmt.accent_color
    lh, gap = fmt.line_height, fmt.section_gap
    base = f"""
body {{ font-family: {family}; font-size: {body}pt; line-height: {lh}; color: #1A1F2B; }}
h1 {{ font-size: {name}pt; margin: 0; color: {accent}; }}
p {{ margin: 0 0 2pt 0; }}
p.headline {{ font-size: {body + 1}pt; margin: 1pt 0 1pt 0; }}
p.contact {{ font-size: {max(MIN_BODY_SIZE, body - 0.5)}pt; color: #3D4657; margin: 0 0 2pt 0; }}
h3 {{ font-size: {body + 0.5}pt; margin: {gap * 0.45}pt 0 0 0; font-weight: bold; }}
p.meta {{ color: #3D4657; margin: 0 0 1pt 0; }}
ul {{ margin: 1pt 0 0 0; padding-left: {body * 1.25}pt; }}
li {{ margin: 0 0 1.5pt 0; }}
"""
    if fmt.template == "modern":
        base += f"""
h2 {{ font-size: {head}pt; color: {accent}; margin: {gap}pt 0 2pt 0; letter-spacing: 0.6pt; }}
"""
    elif fmt.template == "technical":
        base += f"""
h1 {{ font-size: {name}pt; color: #111827; }}
h2 {{ font-size: {head}pt; color: {accent}; margin: {gap}pt 0 2pt 0; padding: 0 0 1pt 0;
      border-bottom: 1.2pt solid {accent}; text-transform: uppercase; }}
p.inline {{ margin: 0 0 1pt 0; }}
"""
    else:  # classic and original
        base += f"""
h2 {{ font-size: {head}pt; color: {accent}; margin: {gap}pt 0 3pt 0; padding: 0 0 1.5pt 0;
      border-bottom: 0.6pt solid #8A94A6; text-transform: uppercase; letter-spacing: 0.4pt; }}
"""
    return base


def _html(content: CVDocument, fmt: FormatSettings, breaks: set[str]) -> str:
    out = [f"<h1>{escape(content.contact.name or 'Curriculum Vitae')}</h1>"]
    if content.contact.headline:
        out.append(f'<p class="headline">{escape(content.contact.headline)}</p>')
    if line := contact_line(content):
        out.append(f'<p class="contact">{escape(line)}</p>')
    for s in visible_sections(content):
        style = ' style="page-break-before: always"' if f"s-{s.id}" in breaks else ""
        title = s.title if fmt.template == "modern" else s.title.upper()  # Story ignores text-transform
        out.append(f'<h2 id="s-{s.id}"{style}>{escape(title)}</h2>')
        for e in s.entries:
            items = [escape(i.text) for i in e.items]
            brk = ' style="page-break-before: always"'
            if s.kind in INLINE_KINDS:
                label = f"<b>{escape(e.title)}:</b> " if e.title else ""
                sep = " · " if fmt.template in ("modern", "technical") else ", "
                out.append(f'<p class="inline">{label}{sep.join(items)}</p>')
                continue
            if e.title or e.organization:
                style = ' style="page-break-before: always"' if f"e-{e.id}" in breaks else ""
                head = escape(e.title) + (f" — {escape(e.organization)}" if e.title and e.organization
                                          else escape(e.organization))
                out.append(f'<h3 id="e-{e.id}"{style}>{head}</h3>')
            meta = "  ·  ".join(escape(x) for x in (e.location, e.dates) if x)
            if meta:
                out.append(f'<p class="meta">{meta}</p>')
            if s.kind in ("summary", "references", "custom") and len(items) <= 2:
                out += [f"<p>{t}</p>" for t in items]
            elif items:
                out.append("<ul>" + "".join(f'<li id="i-{i.id}"{brk if f"i-{i.id}" in breaks else ""}>'
                                            f"{escape(i.text)}</li>" for i in e.items) + "</ul>")
    return "<body>" + "".join(out) + "</body>"


@dataclass
class _Layout:
    data: bytes
    pages: int
    stranded: list[str]  # headings left alone at the bottom of a page, and bullets split over two pages


def _layout(content: CVDocument, fmt: FormatSettings, breaks: set[str]) -> _Layout:
    width, height = PAGE_SIZES[fmt.page]
    page = fitz.Rect(0, 0, width, height)
    bottom_room = max(fmt.margin, 40.0)  # room for the page number
    area = fitz.Rect(fmt.margin, fmt.margin * 0.85, width - fmt.margin, height - bottom_room)
    story = fitz.Story(html=_html(content, fmt, breaks), user_css=_css(fmt))
    buf = io.BytesIO()
    writer = fitz.DocumentWriter(buf)
    more, n = True, 0
    while more:
        n += 1
        if n > 12:
            raise RenderError("The CV is too long to lay out (more than 12 pages).")
        device = writer.begin_page(page)
        more, _ = story.place(area)
        story.draw(device)
        writer.end_page()
    writer.close()
    stranded: list[str] = []
    data = buf.getvalue()
    if n > 1:
        stranded += [hid for hid in _split_bullets(data, content) if hid not in breaks]
    return _Layout(data, n, stranded)


def _split_bullets(data: bytes, content: CVDocument) -> list[str]:
    """Read the PDF back: bullets that continue onto the next page, and headings left at the bottom of a page."""
    sections = visible_sections(content)
    items = [(i.id, _norm(i.text)) for s in sections for e in s.entries for i in e.items]
    heads: dict[str, str] = {}
    for s in sections:
        heads[_norm(s.title)] = f"s-{s.id}"
        for e in s.entries:
            if s.kind in INLINE_KINDS or not (e.title or e.organization):
                continue
            header = e.title + (f" — {e.organization}" if e.title and e.organization else e.organization)
            heads[_norm(header)] = f"e-{e.id}"
            if meta := "  ·  ".join(x for x in (e.location, e.dates) if x):
                heads[_norm(meta)] = f"e-{e.id}"
    found = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        pages = list(doc)
        for number, page in enumerate(pages):
            lines = [_norm(text) for _, text in sorted(
                (ln["bbox"][1], "".join(sp["text"] for sp in ln["spans"]))
                for b in page.get_text("dict")["blocks"] for ln in b.get("lines", []))]
            lines = [ln for ln in lines if ln and not _FOOTER_LINE.match(ln)]
            if not lines:
                continue
            if number + 1 == len(pages):
                break
            last = lines[-1]
            if last in heads:  # a heading with nothing under it
                found.append(heads[last])
            elif last[:1] in "•·-*":  # the page ends with the first line(s) of a bullet
                start = last.lstrip("•·-* ")
                found += [f"i-{item_id}" for item_id, text in items[:]
                          if text.startswith(start) and len(text) > len(start) + 1][:1]
    return list(dict.fromkeys(found))


def _number_pages(data: bytes, fmt: FormatSettings, name: str) -> bytes:
    doc = fitz.open(stream=data, filetype="pdf")
    total = doc.page_count
    if total > 1:
        for i, page in enumerate(doc, start=1):
            label = f"{name} — page {i} of {total}" if name else f"Page {i} of {total}"
            size = 8
            width = fitz.get_text_length(label, fontname="helv", fontsize=size)
            y = page.rect.height - max(fmt.margin, 40.0) / 2
            page.insert_text(((page.rect.width - width) / 2, y), label, fontname="helv", fontsize=size,
                             color=(0.42, 0.45, 0.5))
    doc.set_metadata({"title": f"{name} — CV" if name else "CV", "author": name or "", "creator": "CareerMatch AI",
                      "producer": "CareerMatch AI"})
    out = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return out


def _fit_steps(fmt: FormatSettings):
    """Tighter but still readable settings, in the order the user notices them least."""
    yield fmt
    for gap, lh in ((6, 1.25), (4, 1.2), (3, 1.15)):
        fmt = fmt.model_copy(update={"section_gap": min(fmt.section_gap, gap), "line_height": min(fmt.line_height, lh)})
        yield fmt
    for margin in (44, MIN_MARGIN):
        fmt = fmt.model_copy(update={"margin": min(fmt.margin, margin)})
        yield fmt
    while fmt.body_size > MIN_BODY_SIZE:
        size = max(MIN_BODY_SIZE, round(fmt.body_size - 0.5, 1))
        fmt = fmt.model_copy(update={"body_size": size, "heading_size": max(9, min(fmt.heading_size, size + 1)),
                                     "name_size": max(12, min(fmt.name_size, size + 7))})
        yield fmt


@dataclass
class PDFResult:
    data: bytes
    pages: int
    format: FormatSettings  # the settings actually used (may be tighter when fitting to one page)
    fits_one_page: bool


def render_pdf(content: CVDocument, fmt: FormatSettings, *, one_page: bool = False) -> PDFResult:
    if not visible_sections(content) and not content.contact.name:
        raise RenderError("The CV is empty — there is nothing to generate.")
    candidates = _fit_steps(fmt) if one_page else iter([fmt])
    best: tuple[_Layout, FormatSettings] | None = None
    for candidate in candidates:
        breaks: set[str] = set()
        layout = _layout(content, candidate, breaks)
        for _ in range(_MAX_LAYOUT_PASSES):  # keep headings with the text that follows them
            if not layout.stranded:
                break
            breaks |= set(layout.stranded[:1])  # one at a time: each fix moves everything after it
            layout = _layout(content, candidate, breaks)
        best = (layout, candidate)
        if not one_page or layout.pages == 1:
            break
    layout, used = best
    data = _number_pages(layout.data, used, content.contact.name)
    return PDFResult(data, layout.pages, used, layout.pages == 1)


# --- DOCX -----------------------------------------------------------------------------------

def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _page_number_field(paragraph) -> None:
    for code in ("PAGE",):
        run = paragraph.add_run()
        begin, instr, end = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr.set(qn("xml:space"), "preserve")
        instr.text = code
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(instr)
        run._r.append(end)


def _bottom_border(paragraph, color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", color.lstrip("#"))):
        bottom.set(qn(k), v)
    borders.append(bottom)
    p_pr.append(borders)


def render_docx(content: CVDocument, fmt: FormatSettings) -> bytes:
    """An editable Word document: real headings, paragraphs and bullet lists — no text boxes or tables."""
    if not visible_sections(content) and not content.contact.name:
        raise RenderError("The CV is empty — there is nothing to generate.")
    doc = Document()
    section = doc.sections[0]
    width, height = PAGE_SIZES[fmt.page]
    section.page_width, section.page_height = Pt(width), Pt(height)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Pt(fmt.margin))
    font = _DOCX_FONT[fmt.font_family]
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = font, Pt(fmt.body_size)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = fmt.line_height * 0.9
    for name, size in (("Title", fmt.name_size), ("Heading 1", fmt.heading_size), ("Heading 2", fmt.body_size + 0.5)):
        style = doc.styles[name]
        style.font.name, style.font.size, style.font.bold = font, Pt(size), True
        style.font.color.rgb = _rgb(fmt.accent_color if name != "Heading 2" else "#1A1F2B")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(fmt.section_gap if name == "Heading 1" else 4 if name == "Heading 2" else 0)
        style.paragraph_format.space_after = Pt(2)
    doc.styles["Title"].font.italic = False
    doc.styles["Heading 2"].font.italic = False
    doc.add_paragraph(content.contact.name or "Curriculum Vitae", style="Title")
    if content.contact.headline:
        p = doc.add_paragraph()
        p.add_run(content.contact.headline).font.size = Pt(fmt.body_size + 1)
    if line := contact_line(content):
        doc.add_paragraph(line)
    for s in visible_sections(content):
        heading = doc.add_paragraph(s.title.upper() if fmt.template != "modern" else s.title, style="Heading 1")
        if fmt.template in ("classic", "original", "technical"):
            _bottom_border(heading, fmt.accent_color if fmt.template == "technical" else "#8A94A6")
        for e in s.entries:
            if s.kind in INLINE_KINDS:
                p = doc.add_paragraph()
                if e.title:
                    p.add_run(f"{e.title}: ").bold = True
                p.add_run((" · " if fmt.template in ("modern", "technical") else ", ").join(i.text for i in e.items))
                continue
            if e.title or e.organization:
                doc.add_paragraph(e.title + (f" — {e.organization}" if e.title and e.organization else e.organization),
                                  style="Heading 2")
            meta = "  ·  ".join(x for x in (e.location, e.dates) if x)
            if meta:
                p = doc.add_paragraph(meta)
                p.paragraph_format.keep_with_next = True
            if s.kind in ("summary", "references", "custom") and len(e.items) <= 2:
                for i in e.items:
                    doc.add_paragraph(i.text)
            else:
                for i in e.items:
                    doc.add_paragraph(i.text, style="List Bullet")
    footer = section.footer.paragraphs[0]
    footer.add_run(f"{content.contact.name} — page " if content.contact.name else "Page ").font.size = Pt(8)
    _page_number_field(footer)
    doc.core_properties.title = f"{content.contact.name} — CV" if content.contact.name else "CV"
    doc.core_properties.author = content.contact.name or ""
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- validation ------------------------------------------------------------------------------

@dataclass
class Validation:
    kind: str  # pdf / docx
    ok: bool
    pages: int | None
    checks: list[tuple[str, bool, str]] = field(default_factory=list)  # (label, passed, detail)

    @property
    def problems(self) -> list[str]:
        return [f"{label}: {detail}" for label, passed, detail in self.checks if not passed]

    def as_dict(self) -> dict:
        return {"kind": self.kind, "ok": self.ok, "pages": self.pages,
                "checks": [{"label": a, "passed": b, "detail": c} for a, b, c in self.checks]}


def _missing(expected: list[str], text: str) -> list[str]:
    haystack = _norm(text)
    return [line for line in expected if _norm(line) not in haystack]


def _in_order(expected: list[str], text: str) -> bool:
    haystack, pos = _norm(text), 0
    for line in expected:
        found = haystack.find(_norm(line), pos)
        if found < 0:
            return False
        pos = found
    return True


def validate_pdf(data: bytes, content: CVDocument) -> Validation:
    checks: list[tuple[str, bool, str]] = []
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception:
        return Validation("pdf", False, None, [("Opens", False, "The file could not be opened as a PDF.")])
    with doc:
        checks.append(("Opens", True, f"Valid PDF, {doc.page_count} page(s)"))
        text = _FOOTER_LINE.sub("", "\n".join(page.get_text() for page in doc))  # page numbers split lines
        expected = expected_lines(content)
        missing = _missing(expected, text)
        checks.append(("All content present", not missing,
                       "Every visible line was found in the file" if not missing
                       else f"{len(missing)} line(s) missing, e.g. “{missing[0][:80]}”"))
        titles = [s.title for s in visible_sections(content)]
        checks.append(("Reading order", _in_order(titles, text), "Sections read in the right order"
                       if _in_order(titles, text) else "Sections are read out of order"))
        clipped, overlaps = 0, 0
        for page in doc:
            rect = page.rect
            boxes = []
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    box = fitz.Rect(line["bbox"])
                    if not rect.contains(box):
                        clipped += 1
                    boxes.append(box)
            for i, a in enumerate(boxes):
                for b in boxes[i + 1:]:
                    inter = a & b
                    if not inter.is_empty and inter.get_area() > 0.3 * min(a.get_area(), b.get_area()) > 0:
                        overlaps += 1
        checks.append(("Nothing clipped", clipped == 0, "All text is inside the page" if not clipped
                       else f"{clipped} line(s) run off the page"))
        checks.append(("No overlapping text", overlaps == 0, "No lines overlap" if not overlaps
                       else f"{overlaps} overlapping line(s)"))
        broken = "�" in text
        checks.append(("Characters", not broken, "All characters display correctly" if not broken
                       else "Some characters could not be displayed"))
        selectable = bool(text.strip())
        checks.append(("Selectable text", selectable, "Text can be selected and searched" if selectable
                       else "No selectable text"))
        pages = doc.page_count
    return Validation("pdf", all(c[1] for c in checks), pages, checks)


def validate_docx(data: bytes, content: CVDocument) -> Validation:
    checks: list[tuple[str, bool, str]] = []
    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        return Validation("docx", False, None, [("Opens", False, "The file could not be opened as a Word document.")])
    checks.append(("Opens", True, "Valid Word document"))
    text = "\n".join(p.text for p in doc.paragraphs)
    missing = _missing(expected_lines(content), text)
    checks.append(("All content present", not missing, "Every visible line was found in the file" if not missing
                   else f"{len(missing)} line(s) missing, e.g. “{missing[0][:80]}”"))
    titles = [s.title for s in visible_sections(content)]
    checks.append(("Reading order", _in_order(titles, text), "Sections read in the right order"
                   if _in_order(titles, text) else "Sections are read out of order"))
    body = doc.element.body.xml
    simple = "txbxContent" not in body and not doc.tables
    checks.append(("Easy to edit", simple, "Plain paragraphs and bullet lists, no text boxes or tables" if simple
                   else "Contains text boxes or tables"))
    return Validation("docx", all(c[1] for c in checks), None, checks)


def preview_png(pdf: bytes, *, dpi: int = 100, max_pages: int = 4) -> list[bytes]:
    with fitz.open(stream=pdf, filetype="pdf") as doc:
        return [page.get_pixmap(dpi=dpi).tobytes("png") for page in list(doc)[:max_pages]]
