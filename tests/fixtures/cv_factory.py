"""Generate synthetic CV files in memory. All content is fictional."""

from __future__ import annotations

import io
import zipfile

import pymupdf as fitz
from docx import Document

SAMPLE_CV_LINES = [
    "Alex Example",
    "alex.example@example.com | +00 000 000 000 | Example City",
    "",
    "PROFESSIONAL SUMMARY",
    "Quality engineer with eight years of experience in manufacturing process improvement,",
    "statistical process control and supplier quality management.",
    "",
    "EXPERIENCE",
    "Senior Quality Engineer — Example Manufacturing Ltd — Jan 2019 to Present",
    "• Led root-cause analysis using 8D and Ishikawa methods across three production lines.",
    "• Introduced SPC dashboards in Python and Power BI for critical dimensions.",
    "Quality Engineer — Sample Components Co — Jun 2015 to Dec 2018",
    "• Managed supplier PPAP submissions and incoming inspection plans.",
    "",
    "EDUCATION",
    "BSc Mechanical Engineering — Example University — 2015",
    "",
    "SKILLS",
    "Six Sigma Green Belt, ISO 9001, FMEA, SPC, Python, SQL, Minitab",
]
SAMPLE_CV_TEXT = "\n".join(SAMPLE_CV_LINES)


def make_text_pdf(lines: list[str] | None = None, pages: int = 1) -> bytes:
    doc = fitz.open()
    for _ in range(pages):
        page = doc.new_page()
        y = 60
        for line in lines if lines is not None else SAMPLE_CV_LINES:
            page.insert_text((50, y), line, fontsize=10)
            y += 14
    data = doc.tobytes()
    doc.close()
    return data


def _png_bytes() -> bytes:
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 260), False)
    pix.set_rect(pix.irect, (235, 235, 235))
    return pix.tobytes("png")


def make_image_only_pdf(pages: int = 1) -> bytes:
    doc = fitz.open()
    png = _png_bytes()
    for _ in range(pages):
        page = doc.new_page()
        page.insert_image(page.rect, stream=png)
    data = doc.tobytes()
    doc.close()
    return data


def make_mixed_pdf() -> bytes:
    """Page 1 has text, page 2 is an image only."""
    doc = fitz.open()
    page = doc.new_page()
    y = 60
    for line in SAMPLE_CV_LINES:
        page.insert_text((50, y), line, fontsize=10)
        y += 14
    page2 = doc.new_page()
    page2.insert_image(page2.rect, stream=_png_bytes())
    data = doc.tobytes()
    doc.close()
    return data


def make_encrypted_pdf() -> bytes:
    doc = fitz.open()
    doc.new_page().insert_text((50, 60), "Secret CV", fontsize=10)
    data = doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
    doc.close()
    return data


def make_docx(with_table: bool = True, with_header: bool = True, heading_styles: bool = True) -> bytes:
    document = Document()
    if with_header:
        document.sections[0].header.paragraphs[0].text = "alex.example@example.com | +00 000 000 000"
    document.add_paragraph("Alex Example", style="Title" if heading_styles else None)
    for section, body in [
        ("Professional Summary", ["Quality engineer with eight years of experience in process improvement."]),
        ("Experience", ["Senior Quality Engineer — Example Manufacturing Ltd — 2019 to Present"]),
    ]:
        if heading_styles:
            document.add_heading(section, level=1)
        else:
            document.add_paragraph(section.upper())
        for line in body:
            document.add_paragraph(line)
    document.add_paragraph("Led root-cause analysis using 8D methods.", style="List Bullet")
    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Skills"
        table.cell(0, 1).text = "Python, SQL, SPC, FMEA"
        table.cell(1, 0).text = "Languages"
        table.cell(1, 1).text = "English, Arabic"
    document.add_heading("Education", level=1) if heading_styles else document.add_paragraph("EDUCATION")
    document.add_paragraph("BSc Mechanical Engineering — Example University — 2015")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_empty_docx() -> bytes:
    buf = io.BytesIO()
    Document().save(buf)
    return buf.getvalue()


def make_zip_without_document() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "not a word file")
    return buf.getvalue()
